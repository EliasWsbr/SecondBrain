#!/usr/bin/env python3
"""Baut aus einer Wochennotiz (2026-Www.md) den ausgefuellten Ausbildungsnachweis
als .docx und .pdf, auf Basis der verbindlichen Vorlage.

Aufruf:
    python3 build_berichtsheft.py <woche.md> <vorlage.docx> <outdir> [--datum TT.MM.JJJJ]
                                  [--signatur unterschrift.png]
"""
import argparse
import copy
import datetime as dt
import re
import subprocess
import sys
from pathlib import Path

from docx import Document
from docx.shared import Cm, Pt
from docx.enum.table import WD_ALIGN_VERTICAL, WD_ROW_HEIGHT_RULE
from docx.enum.text import WD_LINE_SPACING

NAME = "{{VOLLER_NAME}}"
ORT = "{{ORT}}"


# ---------------------------------------------------------------- Markdown ---

def parse_week(md_path: Path) -> dict:
    text = md_path.read_text(encoding="utf-8")

    fm = {}
    m = re.match(r"^---\n(.*?)\n---\n", text, re.S)
    if not m:
        sys.exit(f"{md_path.name}: kein Frontmatter gefunden")
    for line in m.group(1).splitlines():
        if ":" in line:
            k, v = line.split(":", 1)
            fm[k.strip()] = v.strip()
    body = text[m.end():]

    # Abschnitte auf ## splitten
    sections, current = {}, None
    for line in body.splitlines():
        h = re.match(r"^##\s+(.*)", line)
        if h:
            current = h.group(1).strip()
            sections[current] = []
        elif current:
            sections[current].append(line)

    def bullets(name):
        out = []
        for line in sections.get(name, []):
            b = re.match(r"^-\s+(.+)$", line.strip())
            if b and b.group(1).strip() not in ("", "_PDF folgt nach Einreichung_"):
                out.append(b.group(1).strip())
        return out

    def hours(name):
        for line in sections.get(name, []):
            h = re.search(r"\*\((\d+[,.]\d)\s*Stunden\)\*", line)
            if h:
                return h.group(1).replace(".", ",")
        return ""

    # Berufsschul-Tabelle: Zeilen mit Inhalt jenseits der Kopf- und Trennzeile
    schule = []
    for line in sections.get("Berufsschule", []):
        s = line.strip()
        if not s.startswith("|") or set(s) <= set("|- "):
            continue
        cells = [c.strip() for c in s.strip("|").split("|")]
        if cells and cells[0].lower() in ("lernfeld / fach", "lernfeld/fach"):
            continue
        if any(cells):
            schule.append(": ".join(c for c in cells if c))

    def de(iso):
        return dt.date.fromisoformat(iso).strftime("%d.%m.%Y")

    return {
        "woche": fm.get("woche", ""),
        "kw": fm.get("woche", "")[-2:],
        "jahr": fm.get("woche", "")[:4],
        "von": de(fm["von"]),
        "bis": de(fm["bis"]),
        "abteilung": fm.get("abteilung", ""),
        "jahr_ausb": f'{fm.get("ausbildungsjahr", "1")}. Ausbildungsjahr',
        "status": fm.get("status", ""),
        "taetigkeiten": bullets("Betriebliche Tätigkeiten"),
        "std_taetigkeiten": hours("Betriebliche Tätigkeiten"),
        "unterweisungen": bullets("Unterweisungen / Schulungen"),
        "std_unterweisungen": hours("Unterweisungen / Schulungen"),
        "schule": schule,
        "std_schule": hours("Berufsschule"),
    }


# -------------------------------------------------------------------- DOCX ---

def cell_text(cell, value, bold=False):
    """Setzt schlichten Text in eine Zelle, Formatvorlage der Vorlage bleibt."""
    p = cell.paragraphs[0]
    for r in list(p.runs):
        r._element.getparent().remove(r._element)
    run = p.add_run(value)
    run.bold = bold
    return p


def cell_bullets(cell, items, spacing=1.15, font_pt=11.0):
    """Schreibt Stichpunkte mit Spiegelstrich und haengendem Einzug."""
    base = cell.paragraphs[0]
    for r in list(base.runs):
        r._element.getparent().remove(r._element)

    def style(p, text):
        pf = p.paragraph_format
        pf.left_indent = Cm(0.85)
        pf.first_line_indent = Cm(-0.4)
        pf.space_before = Pt(0)
        pf.space_after = Pt(0)
        pf.line_spacing = spacing
        run = p.add_run("-\t" + text)
        run.font.size = Pt(font_pt)

    if not items:
        return
    style(base, items[0])
    for item in items[1:]:
        new = copy.deepcopy(base._element)
        base._element.addnext(new) if False else cell._element.append(new)
        p = cell.paragraphs[-1]
        for r in list(p.runs):
            r._element.getparent().remove(r._element)
        style(p, item)


def build(data: dict, template: Path, out_docx: Path, datum: str,
          spacing=1.15, font_pt=11.0, tight_body=False, signatur=None):
    doc = Document(str(template))

    kopf = doc.tables[0]
    cell_text(kopf.rows[0].cells[1], NAME)
    cell_text(kopf.rows[1].cells[1], data["jahr_ausb"])
    cell_text(kopf.rows[2].cells[1], data["abteilung"])
    cell_text(kopf.rows[3].cells[1], data["von"])
    cell_text(kopf.rows[3].cells[3], data["bis"])

    inhalt = doc.tables[1]
    blocks = [
        (1, data["taetigkeiten"], data["std_taetigkeiten"]),
        (3, data["unterweisungen"], data["std_unterweisungen"]),
        (5, data["schule"], data["std_schule"]),
    ]
    for row_idx, items, std in blocks:
        row = inhalt.rows[row_idx]
        cell_bullets(row.cells[0], items, spacing, font_pt)
        cell_text(row.cells[1], std)
        for c in row.cells:
            c.vertical_alignment = WD_ALIGN_VERTICAL.TOP
        # Die Vorlage setzt 4 cm Mindesthoehe pro Inhaltszeile - im eingereichten
        # Nachweis sind leere Abschnitte aber schmal. Hoehe dem Inhalt ueberlassen.
        row.height = None
        row.height_rule = WD_ROW_HEIGHT_RULE.AUTO

    # "Ort, TT.MM.JJJJ" in den ersten leeren Absatz nach dem Bestaetigungssatz
    for i, p in enumerate(doc.paragraphs):
        if p.text.startswith("Durch die nachfolgende Unterschrift"):
            rest = doc.paragraphs[i + 1:]
            for j, q in enumerate(rest):
                if not q.text.strip():
                    q.add_run(f"{ORT}, {datum}")
                    # Unterschrift wie in den gegengezeichneten Nachweisen:
                    # eigene Zeile unter Ort und Datum, direkt auf der Linie
                    # "Datum, Unterschrift Auszubildende/r".
                    if signatur and j + 1 < len(rest) and not rest[j + 1].text.strip():
                        s = rest[j + 1]
                        s.paragraph_format.space_before = Pt(4)
                        s.paragraph_format.space_after = Pt(0)
                        s.paragraph_format.line_spacing = 1.0
                        s.add_run().add_picture(str(signatur), width=Cm(4.4))
                        # Linienzeile mit fester, kleiner Hoehe, damit die
                        # Unterschrift auf der Linie aufsitzt statt darueber
                        # zu schweben.
                        if j + 2 < len(rest):
                            linie = rest[j + 2].paragraph_format
                            linie.space_before = Pt(0)
                            linie.line_spacing_rule = WD_LINE_SPACING.EXACTLY
                            linie.line_spacing = Pt(7)
                    break
            break

    if tight_body:
        # Absaetze ausserhalb der Tabellen enger stellen - spart Zeilen im
        # Unterschriftsblock, ohne den Tabelleninhalt anzutasten.
        for p in doc.paragraphs:
            if p.paragraph_format.line_spacing_rule == WD_LINE_SPACING.EXACTLY:
                continue
            p.paragraph_format.line_spacing = 1.0
            p.paragraph_format.space_after = Pt(0)
        # Ort und Datum mit Abstand zum Bestaetigungssatz, wie im Original.
        for p in doc.paragraphs:
            if p.text.startswith(f"{ORT}, "):
                p.paragraph_format.space_before = Pt(8)

    out_docx.parent.mkdir(parents=True, exist_ok=True)
    doc.save(str(out_docx))


def to_pdf(docx_path: Path, outdir: Path) -> Path:
    subprocess.run(
        ["soffice", "--headless", "--convert-to", "pdf", "--outdir", str(outdir), str(docx_path)],
        check=True, capture_output=True, timeout=180,
    )
    return outdir / (docx_path.stem + ".pdf")


def page_count(pdf: Path) -> int:
    """Seitenzahl ueber pdfinfo, sonst direkt aus der PDF-Struktur gelesen.

    pdfinfo gehoert zu poppler und fehlt unter Windows regelmaessig. Ohne den
    Fallback meldet das Skript dort faelschlich, der Nachweis passe nicht auf
    eine Seite, und bricht mit Code 2 ab.
    """
    try:
        out = subprocess.run(["pdfinfo", str(pdf)],
                             capture_output=True, text=True).stdout
        m = re.search(r"Pages:\s+(\d+)", out)
        if m:
            return int(m.group(1))
    except (FileNotFoundError, OSError):
        pass
    raw = pdf.read_bytes()
    counts = [int(x) for x in re.findall(rb"/Count\s+(\d+)", raw)]
    if counts:
        return max(counts)
    return len(re.findall(rb"/Type\s*/Page[^s]", raw)) or 99


# Es wird IMMER mit den Massen des eingereichten Nachweises gesetzt:
# Arial 11 pt, Zeilenabstand 1,15. Wird es zu lang, werden die Stichpunkte
# gekuerzt - nicht der Satz gequetscht. Das Skript meldet das und bricht ab.
SATZ = (1.15, 11.0, False)


def main():
    ap = argparse.ArgumentParser()
    ap.add_argument("week_md")
    ap.add_argument("template")
    ap.add_argument("outdir")
    ap.add_argument("--datum", default=dt.date.today().strftime("%d.%m.%Y"))
    # Nur der Unterschriftsblock wird enger gesetzt, der Tabelleninhalt nicht.
    # Noetig, wenn LibreOffice mit echtem Arial statt Liberation Sans rendert:
    # die Zeilenhoehe faellt dann groesser aus und der Block rutscht auf Seite 2.
    # Gegenprobe ist die erste eingereichte Woche: sie liegt als PDF auf einer
    # Seite vor und muss mit diesem Schalter auch hier auf eine passen.
    ap.add_argument("--eng", action="store_true",
                    help="Absaetze ausserhalb der Tabellen enger setzen")
    # Bild der eigenen Unterschrift. Liegt bewusst ausserhalb des Vaults,
    # damit es nicht per Git synchronisiert wird.
    ap.add_argument("--signatur", type=Path,
                    help="PNG der Unterschrift des Auszubildenden")
    a = ap.parse_args()

    data = parse_week(Path(a.week_md))
    outdir = Path(a.outdir)
    stem = f'{NAME}_Berichtsheft_KW{data["kw"]}_{data["jahr"]}'
    out_docx = outdir / f"{stem}.docx"

    satz = SATZ[:2] + (a.eng,)
    build(data, Path(a.template), out_docx, a.datum, *satz, signatur=a.signatur)
    pdf = to_pdf(out_docx, outdir)
    seiten = page_count(pdf)

    total = 0.0
    for key in ("std_taetigkeiten", "std_unterweisungen", "std_schule"):
        if data[key]:
            total += float(data[key].replace(",", "."))
    warn = "" if abs(total - 40.0) < 0.05 else f"  ACHTUNG: Summe {total:.1f} h, erwartet 40,0 h"

    print(f'{data["woche"]}  status={data["status"]}  '
          f'{len(data["taetigkeiten"])} Stichpunkte  '
          f'{len(data["unterweisungen"])} Unterweisungen  '
          f'Summe {total:.1f} h{warn}')
    print(f"  {pdf}  Seiten: {seiten}")

    if seiten > 1:
        print("\n  PASST NICHT AUF EINE SEITE.")
        print("  Der Satz wird NICHT enger gestellt - stattdessen die Stichpunkte")
        print("  in der Wochennotiz kuerzen: lange Punkte straffen, eng verwandte")
        print("  zusammenfassen, Fuellwoerter raus. Danach erneut bauen.")
        sys.exit(2)


if __name__ == "__main__":
    main()
